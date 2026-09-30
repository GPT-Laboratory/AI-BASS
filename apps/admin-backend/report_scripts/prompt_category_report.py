from __future__ import annotations

import json
import os
import re
from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from .report_utils import collect_filtered_prompt_logs, resolve_time_range


class PromptCategoryReport:
    script_id = "prompt-category-report"
    report_model = os.getenv("AZURE_REPORT_MODEL", "gpt-4o-mini")
    fixed_personal_category = "Personal prompts"

    def describe(self):
        return {
            "id": self.script_id,
            "name": "Prompt Categories by Company",
            "output_format": "xlsx",
        }

    def run(self, context: dict, params: dict):
        mongo = context["mongo"]
        azure_client = context["azure_client"]
        logger = context["logger"]

        start_time, end_time = resolve_time_range(params)
        dataset = self._collect_dataset(mongo, start_time, end_time, params)
        unique_prompts = list(dataset["unique_prompts"].values())

        logger.info(
            "[report:%s] Collected %s production log rows and %s unique prompts",
            self.script_id,
            dataset["log_count"],
            len(unique_prompts),
        )

        categories = self._generate_categories(azure_client, unique_prompts, logger)
        classifications = self._classify_prompts(azure_client, unique_prompts, categories, logger)
        workbook_bytes = self._build_workbook(dataset, classifications, categories, start_time, end_time)
        filename = self._build_filename(start_time, end_time)

        return {
            "content": workbook_bytes,
            "filename": filename,
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }

    def _collect_dataset(self, mongo, start_time, end_time, params: dict):
        filtered = collect_filtered_prompt_logs(mongo, start_time, end_time, params)
        logs = filtered["logs"]
        company_ids = filtered["company_ids"]
        company_id_by_user = filtered["company_id_by_user"]

        unique_prompts = {}
        company_prompt_counts = defaultdict(Counter)
        company_total_prompts = Counter()

        for log in logs:
            raw_prompt = (log.get("question") or "").strip()
            if not raw_prompt:
                continue

            metrics = log.get("metrics", {}) or {}
            company_id = metrics.get("company_id") or company_id_by_user.get(metrics.get("user_id"))
            if not company_id or company_id not in company_ids:
                continue

            normalized = self._normalize_prompt(raw_prompt)
            if not normalized:
                continue

            prompt_record = unique_prompts.get(normalized)
            if not prompt_record:
                prompt_record = {
                    "id": f"prompt_{len(unique_prompts) + 1}",
                    "normalized": normalized,
                    "sample_text": raw_prompt,
                    "count": 0,
                    "company_counts": Counter(),
                }
                unique_prompts[normalized] = prompt_record

            prompt_record["count"] += 1
            prompt_record["company_counts"][company_id] += 1
            company_prompt_counts[company_id][normalized] += 1
            company_total_prompts[company_id] += 1

        return {
            "company_ids": company_ids,
            "unique_prompts": unique_prompts,
            "company_prompt_counts": company_prompt_counts,
            "company_total_prompts": company_total_prompts,
            "log_count": len(logs),
            "filters": filtered["filters"],
        }

    def _normalize_prompt(self, text: str) -> str:
        normalized = re.sub(r"\s+", " ", (text or "").strip().lower())
        return normalized[:4000]

    def _generate_categories(self, azure_client, unique_prompts: list[dict], logger):
        if not unique_prompts:
            return []

        ranked_prompts = sorted(unique_prompts, key=lambda item: item["count"], reverse=True)
        sample = ranked_prompts[:80]
        prompt_lines = [
            f"{item['id']}: {self._truncate(item['sample_text'], 280)}"
            for item in sample
        ]

        system_prompt = (
            "You design concise business prompt taxonomies. "
            "Return only JSON."
        )
        user_prompt = (
            "Create a compact category list for business-related prompts used by company employees.\n"
            "Rules:\n"
            f"- Do not include the fixed category '{self.fixed_personal_category}'. It already exists separately.\n"
            "- Create 4 to 8 categories total.\n"
            "- Category names must be short English labels.\n"
            "- Categories must be broad enough to cover repeated employee work prompts.\n"
            "- Exclude personal-life prompts from the taxonomy.\n"
            "- Output JSON object with key 'categories', value array of objects with keys 'name' and 'description'.\n\n"
            "Sample prompts:\n"
            + "\n".join(prompt_lines)
        )

        fallback = [
            {"name": "Operations", "description": "Daily work, processes, tasks, instructions, and execution."},
            {"name": "Customer and Sales", "description": "Customers, sales, marketing, offers, communication, and service."},
            {"name": "Strategy and Management", "description": "Planning, decisions, goals, leadership, and management."},
            {"name": "Finance and Reporting", "description": "Costs, pricing, budgeting, profitability, and reporting."},
            {"name": "People and HR", "description": "Employees, recruitment, roles, training, and workplace matters."},
            {"name": "Technology and Data", "description": "Systems, software, automation, data, and technical troubleshooting."},
        ]

        try:
            response = azure_client.chat.completions.create(
                model=self.report_model,
                temperature=0,
                response_format={"type": "json_object"},
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
            )
            content = response.choices[0].message.content or "{}"
            parsed = json.loads(content)
            categories = parsed.get("categories") or []
            cleaned = []
            seen = set()
            for item in categories:
                name = (item.get("name") or "").strip()
                description = (item.get("description") or "").strip()
                if not name:
                    continue
                if name.lower() == self.fixed_personal_category.lower():
                    continue
                if name.lower() in seen:
                    continue
                seen.add(name.lower())
                cleaned.append({"name": name, "description": description})
            if cleaned:
                return cleaned[:8]
        except Exception as exc:
            logger.warning("[report:%s] Category generation failed, using fallback: %s", self.script_id, exc)

        return fallback

    def _classify_prompts(self, azure_client, unique_prompts: list[dict], categories: list[dict], logger):
        if not unique_prompts:
            return {}

        allowed_names = [self.fixed_personal_category] + [item["name"] for item in categories]
        category_descriptions = "\n".join(
            [f"- {self.fixed_personal_category}: Non-work or personal-life prompts."]
            + [f"- {item['name']}: {item.get('description', '')}" for item in categories]
        )

        results = {}
        batch_size = 25
        for batch_start in range(0, len(unique_prompts), batch_size):
            batch = unique_prompts[batch_start:batch_start + batch_size]
            prompt_lines = [
                f"{item['id']}: {self._truncate(item['sample_text'], 500)}"
                for item in batch
            ]
            system_prompt = (
                "You classify employee prompts into one allowed category. "
                "Return only JSON."
            )
            user_prompt = (
                "Classify each prompt into exactly one allowed category.\n"
                "The special category Personal prompts is only for prompts unrelated to company work or the employer's business.\n"
                "If a prompt is ambiguous, prefer a business category over Personal prompts.\n"
                "Use only category names from the allowed list.\n"
                "Output JSON object with key 'classifications', value array of objects with keys "
                "'id', 'category', and optional 'reason'.\n\n"
                "Allowed categories:\n"
                f"{category_descriptions}\n\n"
                "Prompts:\n"
                + "\n".join(prompt_lines)
            )
            try:
                response = azure_client.chat.completions.create(
                    model=self.report_model,
                    temperature=0,
                    response_format={"type": "json_object"},
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt},
                    ],
                )
                content = response.choices[0].message.content or "{}"
                parsed = json.loads(content)
                classifications = parsed.get("classifications") or []
                for item in classifications:
                    prompt_id = item.get("id")
                    category = (item.get("category") or "").strip()
                    reason = (item.get("reason") or "").strip()
                    if prompt_id and category in allowed_names:
                        results[prompt_id] = {"category": category, "reason": reason}
            except Exception as exc:
                logger.warning(
                    "[report:%s] Batch classification failed at %s, using fallback category: %s",
                    self.script_id,
                    batch_start,
                    exc,
                )

            for item in batch:
                if item["id"] not in results:
                    results[item["id"]] = {
                        "category": categories[0]["name"] if categories else self.fixed_personal_category,
                        "reason": "Fallback classification",
                    }

        return results

    def _build_workbook(self, dataset, classifications, categories, start_time, end_time):
        workbook = Workbook()
        summary_ws = workbook.active
        summary_ws.title = "Summary"
        categories_ws = workbook.create_sheet("Categories")
        prompts_ws = workbook.create_sheet("Prompt details")

        company_ids = dataset["company_ids"]
        unique_prompts = dataset["unique_prompts"]

        category_order = [self.fixed_personal_category] + [item["name"] for item in categories]
        if not category_order:
            category_order = [self.fixed_personal_category]

        category_counts_by_company = defaultdict(Counter)
        total_counts_by_category = Counter()

        for prompt in unique_prompts.values():
            classification = classifications.get(prompt["id"], {})
            category = classification.get("category") or self.fixed_personal_category
            for company_id, count in prompt["company_counts"].items():
                category_counts_by_company[company_id][category] += count
                total_counts_by_category[category] += count

        summary_headers = ["Company ID", "Total prompts"] + category_order
        summary_ws.append(summary_headers)
        self._style_header(summary_ws, 1)

        for company_id in sorted(company_ids.keys()):
            row = [company_id, dataset["company_total_prompts"].get(company_id, 0)]
            for category in category_order:
                row.append(category_counts_by_company[company_id].get(category, 0))
            summary_ws.append(row)

        total_row = ["All companies", sum(dataset["company_total_prompts"].values())]
        for category in category_order:
            total_row.append(total_counts_by_category.get(category, 0))
        summary_ws.append(total_row)

        categories_ws.append(["Category", "Description"])
        self._style_header(categories_ws, 1)
        categories_ws.append([self.fixed_personal_category, "Prompts about private life or non-work matters."])
        for item in categories:
            categories_ws.append([item["name"], item.get("description", "")])

        categories_ws.append([])
        categories_ws.append(["Report metadata", "Value"])
        self._style_header(categories_ws, categories_ws.max_row)
        categories_ws.append(["Start time", start_time.isoformat()])
        categories_ws.append(["End time", end_time.isoformat()])
        categories_ws.append(["Environment filter", dataset["filters"].get("company_env") or "all"])
        categories_ws.append(["Company filter", dataset["filters"].get("company_id") or "all"])
        categories_ws.append(["User filter", dataset["filters"].get("user_id") or "all"])
        categories_ws.append(["Unique prompts", len(unique_prompts)])
        categories_ws.append(["Log rows", dataset["log_count"]])

        prompt_headers = ["Prompt ID", "Category", "Reason", "Total count", "Companies", "Prompt text"]
        prompts_ws.append(prompt_headers)
        self._style_header(prompts_ws, 1)

        sorted_prompts = sorted(unique_prompts.values(), key=lambda item: item["count"], reverse=True)
        for prompt in sorted_prompts:
            classification = classifications.get(prompt["id"], {})
            company_counts_text = ", ".join(
                f"{company_id}: {count}"
                for company_id, count in sorted(prompt["company_counts"].items(), key=lambda pair: pair[0])
            )
            prompts_ws.append(
                [
                    prompt["id"],
                    classification.get("category") or self.fixed_personal_category,
                    classification.get("reason", ""),
                    prompt["count"],
                    company_counts_text,
                    prompt["sample_text"],
                ]
            )

        for ws in (summary_ws, categories_ws, prompts_ws):
            self._autosize_columns(ws)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _style_header(self, worksheet, row_number: int):
        for cell in worksheet[row_number]:
            cell.font = Font(bold=True)

    def _autosize_columns(self, worksheet):
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    value = str(cell.value) if cell.value is not None else ""
                except Exception:
                    value = ""
                if len(value) > max_length:
                    max_length = len(value)
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)

    def _truncate(self, value: str, max_len: int) -> str:
        if not value:
            return ""
        return value if len(value) <= max_len else value[: max_len - 3] + "..."

    def _build_filename(self, start_time: datetime, end_time: datetime):
        start_label = start_time.strftime("%Y%m%d")
        end_label = end_time.strftime("%Y%m%d")
        return f"prompt_category_report_{start_label}_{end_label}.xlsx"
